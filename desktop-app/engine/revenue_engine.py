"""محرك الإيرادات: يقرأ ملفات الفواتير من الدرايف، ينقّيها ويصنّفها ويجمعها، ويكتب ملف المتابعة الشهري.

كل الأرقام تُحسب هنا (مو بالذكاء الاصطناعي). المحرك:
- ما يقرأ إلا الأعمدة المحددة في الإعدادات، وأي عمود فيه اسم أو هوية ما يُقرأ أبداً.
- ما يكتب على أي ملف مصدر؛ المخرجات ملفات جديدة فقط.
- يطبع تقدّمه سطر JSON لكل حدث (للمبنى).

الأوامر:
  inspect  --config cfg.json                     يعرض الأوراق والعناوين وعدد الصفوف (بدون بيانات)
  monthly  --config cfg.json --year Y --month M --out DIR
  numcheck --summary summary.json --text FILE    يتأكد إن كل رقم في النص موجود في الملخص
"""
import argparse, csv, datetime as dt, json, os, re, shutil, sys, time
from collections import defaultdict

import openpyxl

# Windows يطبع افتراضياً بترميز cp1252 اللي ما يدعم العربي (حتى في exe المبني بـ PyInstaller)
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass

MONTHS = ['يناير', 'فبراير', 'مارس', 'أبريل', 'مايو', 'يونيو', 'يوليو', 'أغسطس', 'سبتمبر', 'أكتوبر', 'نوفمبر', 'ديسمبر']
REGION, CITY, MUNIS = 'منطقة الرياض', 'مدينة الرياض', 'بلديات المنطقة'
UNCLASSIFIED = 'غير مصنف'


# ---------------------------------------------------------------- events
def emit(ev, **kw):
    kw['ev'] = ev
    kw.setdefault('at', int(time.time() * 1000))
    sys.stdout.write(json.dumps(kw, ensure_ascii=False, default=str) + '\n')
    sys.stdout.flush()


def status(id, text, kind='work', detail='', screen=None):
    emit('status', id=id, kind=kind, text=text, detail=detail, screen=screen or [])


def chat(frm, to, text, kind='finding', **kw):
    emit('chat', **{'from': frm, 'to': to, 'kind': kind, 'text': text, **kw})


# ---------------------------------------------------------------- text helpers
_AR = str.maketrans({'أ': 'ا', 'إ': 'ا', 'آ': 'ا', 'ى': 'ي', 'ة': 'ه', '\xa0': ' ', 'ـ': ''})
_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩٫٬', '0123456789.,')


def norm(s):
    if s is None:
        return ''
    return re.sub(r'\s+', ' ', str(s).translate(_AR)).strip().lower()


def fmt_m(v):
    return f'{v / 1e6:,.1f}'


def to_num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).translate(_DIGITS).replace(',', '').strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def account_key(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).translate(_DIGITS).strip()


def to_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # رقم تاريخ إكسل
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = str(v).translate(_DIGITS).strip()
    for f, n in (('%Y-%m-%d', 10), ('%Y/%m/%d', 10), ('%d/%m/%Y', 10), ('%d-%m-%Y', 10)):
        try:
            return dt.datetime.strptime(s[:n], f).date()
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- config
def load_config(path):
    cfg = json.load(open(path, encoding='utf-8'))
    base = os.path.dirname(os.path.abspath(path))
    root = cfg.get('drive_root') or base

    def p(x):
        if not x:
            return None
        return x if os.path.isabs(x) else os.path.join(root, x)

    cfg['_paths'] = {k: p(v) for k, v in cfg.get('sources', {}).items()}
    return cfg


def is_pii(header, cfg):
    h = norm(header)
    if any(norm(k) == h for k in cfg.get('pii_exempt', ['اسم الحساب', 'اسم البند', 'اسم النشاط', 'اسم البلدية', 'اسم الجهة'])):
        return False
    return any(norm(k) in h for k in cfg.get('pii_keywords', []))


# ---------------------------------------------------------------- reading
def open_rows(path, sheet=None):
    """يرجّع (اسم الورقة، مولّد صفوف) لملف xlsx أو csv بدون تحميله كامل في الذاكرة."""
    if path.lower().endswith(('.csv', '.txt')):
        f = open(path, encoding='utf-8-sig', newline='')
        return os.path.basename(path), (tuple(r) for r in csv.reader(f))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return ws.title, ws.iter_rows(values_only=True)


def find_header(rows, aliases, max_scan=15):
    """يدوّر صف العناوين في أول الصفوف، ويرجّع خريطة: الحقل ← رقم العمود."""
    buf = []
    for i, r in enumerate(rows):
        buf.append(r)
        cells = [norm(c) for c in r]
        found = {}
        for field, names in aliases.items():
            for n in names:
                nn = norm(n)
                if nn in cells:
                    found[field] = cells.index(nn)
                    break
        if len(found) >= 2:
            return i, r, found
        if i >= max_scan:
            break
    raise ValueError('ما لقيت صف العناوين. راجع أسماء الأعمدة في الإعدادات (columns).')


def scan(path, aliases, cfg, label, who, needed):
    sheet_name, rows = open_rows(path)
    hdr_i, hdr, cols = find_header(rows, aliases)
    missing = [f for f in needed if f not in cols]
    if missing:
        raise ValueError(f'{label}: أعمدة ناقصة {missing}. العناوين الموجودة: {[c for c in hdr if c]}')
    for f, ci in cols.items():
        if is_pii(hdr[ci], cfg):
            raise ValueError(f'{label}: العمود «{hdr[ci]}» مربوط بالحقل {f} لكنه يحتوي بيانات شخصية؛ ما راح أقرأه.')
    skipped = [str(c) for c in hdr if c and is_pii(c, cfg)]
    status(who, f'يقرأ {label}', detail=f'ورقة «{sheet_name}»، {len(cols)} عمود مطلوب، {len(skipped)} عمود شخصي متجاهَل',
           screen=[f'فتح {os.path.basename(path)}', f'الأعمدة: {"، ".join(cols)}', f'متجاهَل (بيانات شخصية): {len(skipped)}'])
    n = 0
    for r in rows:
        n += 1
        if n % 200000 == 0:
            status(who, f'يقرأ {label}', detail=f'قرأ {n:,} صف', screen=[f'{label}: {n:,} صف'])
        if not r or all(c is None or c == '' for c in r):
            continue
        yield {f: (r[ci] if ci < len(r) else None) for f, ci in cols.items()}


# ---------------------------------------------------------------- classification
def load_settings(cfg):
    path = cfg['_paths'].get('settings')
    s = cfg.get('settings_sheet', {})
    acc_names = s.get('account', ['رقم الحساب', 'الحساب', 'رمز الحساب'])
    tab_names = s.get('tab', ['التبويب', 'تبويب الامانة'])
    sheet_name, rows = open_rows(path, s.get('sheet'))
    _, hdr, cols = find_header(rows, {'account': acc_names, 'tab': tab_names})
    if 'account' not in cols or 'tab' not in cols:
        raise ValueError('ملف الإعدادات: لازم عمود رقم الحساب وعمود التبويب.')
    m = {}
    for r in rows:
        a = account_key(r[cols['account']] if cols['account'] < len(r) else None)
        t = r[cols['tab']] if cols['tab'] < len(r) else None
        if a and t:
            m[a] = str(t).replace('\xa0', ' ').strip()
    return m


def org_of(muni, cfg):
    n = norm(muni)
    if not n:
        return None
    for mk in cfg.get('city_markers', ['امانه', 'قطاع']):
        if norm(mk) in n:
            return CITY
    return MUNIS


def sector_of(muni, sector, org, cfg):
    if sector:
        return str(sector).replace('\xa0', ' ').strip()
    groups = cfg.get('municipality_groups', {})
    g = groups.get(str(muni).strip()) or groups.get(norm(muni))
    return g or (str(muni).strip() if org == CITY else None)


# ---------------------------------------------------------------- monthly
class Agg:
    def __init__(self):
        self.ytd = defaultdict(lambda: defaultdict(lambda: {'billed': 0.0, 'actual': 0.0, 'prev_years': 0.0}))
        self.month = defaultdict(lambda: defaultdict(float))
        self.sectors = defaultdict(lambda: {'billed': 0.0, 'actual': 0.0, 'prev_years': 0.0})
        self.series = defaultdict(lambda: {'billed': 0.0, 'actual': 0.0, 'prev_years': 0.0})
        self.unclassified = defaultdict(lambda: {'rows': 0, 'amount': 0.0})
        self.no_sector = defaultdict(float)

    def add(self, org, sector, tab, field, v, m=None, month_hit=False):
        for o in (org, REGION):
            self.ytd[o][tab][field] += v
            if month_hit and field == 'actual':
                self.month[o][tab] += v
        if sector:
            self.sectors[(org, sector, tab)][field] += v
        elif org:
            self.no_sector[org] += v
        if m:
            self.series[m][field] += v


def run_monthly(cfg, year, month, out_dir):
    t0 = time.time()
    P = cfg['_paths']
    cols = cfg['columns']
    for k in ('amana_billed', 'settings', 'template'):
        if not P.get(k) or not os.path.exists(P[k]):
            raise FileNotFoundError(f'الملف «{k}» غير موجود: {P.get(k)}')
    period_end = dt.date(year, month, 28) + dt.timedelta(days=4)
    period_end = period_end - dt.timedelta(days=period_end.day)  # آخر يوم في الشهر
    status('angel', 'يوزّع شغل الاستلام', detail=f'متابعة {MONTHS[month - 1]} {year}: الاستلام ثم التصنيف ثم الجهات')
    chat('angel', 'recv', f'ابدأ استلام ملفات {MONTHS[month - 1]} {year} وتنقيتها.', kind='handoff')

    # --- التصنيف
    status('cls', 'يحمّل جدول الحسابات', detail='الإعدادات: رقم الحساب ← التبويب')
    acc_tab = load_settings(cfg)
    status('cls', 'جاهز للتصنيف', detail=f'{len(acc_tab):,} حساب في الجدول', screen=[f'الإعدادات: {len(acc_tab):,} حساب'])

    agg = Agg()
    stats = {}
    cancel_markers = [norm(x) for x in cfg.get('cancel_markers', ['ملغ'])]

    def classify(acc):
        return acc_tab.get(account_key(acc), UNCLASSIFIED)

    # --- مفوتر السنة الحالية (فيه المسدد عن نفس السنة)
    billed_needed = ['account', 'amount', 'issue_date', 'municipality']
    st = stats['amana_billed'] = {'rows': 0, 'used': 0, 'cancelled': 0, 'dup_x': 0, 'x_unresolved': 0, 'other_year': 0, 'no_date': 0}
    invoices, x_rows = set(), []

    def take(rec, st):
        d = to_date(rec.get('issue_date'))
        if not d:
            st['no_date'] += 1
            return
        tab = classify(rec.get('account'))
        org = org_of(rec.get('municipality'), cfg)
        sec = sector_of(rec.get('municipality'), rec.get('sector'), org, cfg)
        amt = to_num(rec.get('amount'))
        if tab == UNCLASSIFIED:
            u = agg.unclassified[account_key(rec.get('account'))]
            u['rows'] += 1
            u['amount'] += amt
        if d.year == year and d <= period_end:
            agg.add(org, sec, tab, 'billed', amt, d.month)
            st['used'] += 1
        else:
            st['other_year'] += 1
        pd_ = to_date(rec.get('pay_date'))
        paid = to_num(rec.get('paid'))
        if paid and pd_ and pd_.year == year and pd_ <= period_end:
            agg.add(org, sec, tab, 'actual', paid, pd_.month, month_hit=pd_.month == month)

    status('recv', 'يستلم ملف المفوتر', detail=os.path.basename(P['amana_billed']))
    for rec in scan(P['amana_billed'], cols, cfg, 'مفوتر السنة', 'recv', billed_needed):
        st['rows'] += 1
        if 'status' in rec and any(m in norm(rec['status']) for m in cancel_markers):
            st['cancelled'] += 1
            continue
        inv = str(rec.get('invoice') or '').strip()
        if inv[-1:] in ('x', 'X'):
            x_rows.append(rec)
            continue
        if inv:
            invoices.add(inv)
        take(rec, st)
    for rec in x_rows:  # الفاتورة اللي تنتهي بـ x تكرار، إلا إذا ما لقينا أصلها
        if str(rec['invoice']).strip()[:-1] in invoices:
            st['dup_x'] += 1
        else:
            st['x_unresolved'] += 1
            take(rec, st)
    invoices.clear()
    status('recv', 'خلّص استلام المفوتر', 'done', detail=f"{st['rows']:,} صف، استُبعد {st['cancelled']:,} ملغي و{st['dup_x']:,} مكرر",
           screen=[f"صفوف: {st['rows']:,}", f"ملغي: {st['cancelled']:,}", f"مكرر x: {st['dup_x']:,}", f"x بدون أصل: {st['x_unresolved']:,}"])
    chat('recv', 'cls', f"سلّمت المفوتر منقّى: {st['used']:,} صف مستخدم، واستبعدت {st['cancelled']:,} ملغي و{st['dup_x']:,} مكرر.", kind='handoff')
    if st['x_unresolved']:
        chat('cls', 'angel', f"{st['x_unresolved']:,} فاتورة تنتهي بـ x وما لقيت أصلها؛ أبقيتها في المجاميع وأرفعها لك.", kind='finding')

    # --- المسدد عن سنوات سابقة
    if P.get('amana_paid_prev') and os.path.exists(P['amana_paid_prev']):
        sp = stats['amana_paid_prev'] = {'rows': 0, 'used': 0, 'cancelled': 0, 'no_date': 0}
        for rec in scan(P['amana_paid_prev'], cols, cfg, 'المسدد عن سنوات سابقة', 'recv', ['account', 'municipality']):
            sp['rows'] += 1
            if 'status' in rec and any(m in norm(rec['status']) for m in cancel_markers):
                sp['cancelled'] += 1
                continue
            d = to_date(rec.get('pay_date')) or to_date(rec.get('issue_date'))
            if not d:
                sp['no_date'] += 1
                continue
            if d.year != year or d > period_end:
                continue
            v = to_num(rec.get('paid')) or to_num(rec.get('amount'))
            tab = classify(rec.get('account'))
            org = org_of(rec.get('municipality'), cfg)
            agg.add(org, sector_of(rec.get('municipality'), rec.get('sector'), org, cfg), tab, 'prev_years', v, d.month)
            sp['used'] += 1
        status('recv', 'خلّص المسدد عن سنوات سابقة', 'done', detail=f"{sp['used']:,} صف مستخدم من {sp['rows']:,}")

    unc_total = sum(u['amount'] for u in agg.unclassified.values())
    status('cls', 'خلّص التصنيف', 'done', detail=f'{len(agg.unclassified)} حساب غير مصنف بمبلغ {fmt_m(unc_total)} مليون',
           screen=[f'{t}: {fmt_m(v["billed"])} مليون مفوتر' for t, v in sorted(agg.ytd[REGION].items())][:8])
    if agg.unclassified:
        top = sorted(agg.unclassified.items(), key=lambda kv: -kv[1]['amount'])[:5]
        chat('cls', 'angel', 'حسابات غير مصنفة (أعلى خمسة): ' + '، '.join(f'{a} ({fmt_m(u["amount"])} مليون)' for a, u in top), kind='finding')
    chat('cls', 'ent', 'صنّفت البنود برقم الحساب، وسلّمتك للتوزيع على الجهات.', kind='handoff')
    status('ent', 'خلّص التوزيع على الجهات', 'done', detail=f'{len(agg.sectors)} مجموعة (جهة × قطاع × تبويب)')

    # --- تحديث القالب
    status('neves', 'يحدّث ملف المتابعة', detail='نسخة جديدة من القالب، بدون لمس المعادلات')
    os.makedirs(out_dir, exist_ok=True)
    out_xlsx = unique_path(os.path.join(out_dir, f'ملف_المتابعة_{year}-{month:02d}.xlsx'))
    tpl = update_template(P['template'], out_xlsx, agg, year, month, cfg)
    status('neves', 'كتب ملف المتابعة', 'done', detail=os.path.basename(out_xlsx),
           screen=[f'خانات محدّثة: {tpl["updated_cells"]}', f'معادلات تُركت: {tpl["skipped_formula"]}', f'مفاتيح ناقصة: {len(tpl["unmatched"])}'])

    # --- الملخص + فحوص Salah
    summary = build_summary(agg, stats, tpl, year, month, unc_total)
    summary['checks'] = salah_checks(summary, agg)
    summary['outputs'] = {'xlsx': out_xlsx}
    summary['elapsed_sec'] = round(time.time() - t0, 1)
    sp = os.path.join(out_dir, 'summary.json')
    json.dump(summary, open(sp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    emit('result', summary=sp, xlsx=out_xlsx)
    return summary


def unique_path(p):
    if not os.path.exists(p):
        return p
    b, e = os.path.splitext(p)
    i = 2
    while os.path.exists(f'{b}_{i}{e}'):
        i += 1
    return f'{b}_{i}{e}'


# ---------------------------------------------------------------- template
def is_formula(c):
    return isinstance(c.value, str) and c.value.startswith('=')


def header_map(ws):
    return {norm(c.value): c.column for c in ws[1] if c.value is not None}


def update_template(src, dst, agg, year, month, cfg):
    """ينسخ القالب لملف جديد ويحدّث الخانات الثابتة فقط (أي خانة فيها معادلة تبقى كما هي)."""
    if os.path.abspath(src) == os.path.abspath(dst):
        raise ValueError('ممنوع الكتابة على ملف المصدر')
    wb = openpyxl.load_workbook(src)  # keep_links افتراضياً
    rep = {'_src': src, 'updated_cells': 0, 'skipped_formula': 0, 'unmatched': [], 'warnings': [], 'prev_month_diffs': [], 'sheets': []}
    for ws in wb.worksheets:
        if getattr(ws, '_charts', None) or getattr(ws, '_images', None):
            rep['warnings'].append(f'الورقة «{ws.title}» فيها رسوم أو صور؛ تأكد منها بعد الفتح.')

    def setv(cell, v):
        if is_formula(cell):
            rep['skipped_formula'] += 1
            return
        cell.value = round(v, 2)
        rep['updated_cells'] += 1

    def sheet(name):
        for ws in wb.worksheets:
            if norm(ws.title) == norm(name):
                return ws
        rep['warnings'].append(f'ما لقيت الورقة «{name}» في القالب.')
        return None

    S = cfg.get('template_sheets', {})
    tab_alias = {norm(k): v for k, v in cfg.get('tab_aliases', {}).items()}

    def tab_key(t):
        n = norm(t)
        return norm(tab_alias.get(n, t))

    # 1) الإيراد حتى تاريخه: (الجهة، التبويب) ← مفوتر/فعلي/محصل سنوات سابقة
    ws = sheet(S.get('ytd', 'الايراد حتى تاريخه'))
    if ws is not None:
        h = header_map(ws)
        c_org, c_tab = h.get(norm('التنظيم الايداري')), h.get(norm('التبويب'))
        fields = {'billed': h.get(norm('مفوتر')), 'actual': h.get(norm('فعلي')), 'prev_years': h.get(norm('محصل  سنوات سابقة'))}
        idx = {(norm(o), tab_key(t)): v for o, tabs in agg.ytd.items() for t, v in tabs.items()}
        seen = set()
        for r in range(2, ws.max_row + 1):
            o, t = ws.cell(r, c_org).value, ws.cell(r, c_tab).value
            if not o or not t:
                continue
            k = (norm(o), tab_key(t))
            v = idx.get(k)
            if v is None:
                if not any(is_formula(ws.cell(r, c)) for c in fields.values() if c):
                    rep['unmatched'].append(f'{ws.title}: {o} / {t}')
                continue
            seen.add(k)
            for f, c in fields.items():
                if c:
                    setv(ws.cell(r, c), v[f])
        for k in idx:
            if k not in seen and k[0] != norm(MUNIS) and k[1] != norm(UNCLASSIFIED):
                rep['unmatched'].append(f'{ws.title}: ما له صف في القالب ← {k[0]} / {k[1]}')
        rep['sheets'].append(ws.title)

    # 2) الإيراد الشهر الحالي: اسم الشهر + فعلي الشهر
    ws = sheet(S.get('month', 'الايراد الشهر الحالي'))
    if ws is not None:
        h = header_map(ws)
        c_m, c_org, c_tab, c_act = h.get(norm('الشهر')), h.get(norm('التنظيم الايداري')), h.get(norm('التبويب')), h.get(norm('فعلي'))
        idx = {(norm(o), tab_key(t)): v for o, tabs in agg.month.items() for t, v in tabs.items()}
        old_month = None
        for r in range(2, ws.max_row + 1):
            o, t = ws.cell(r, c_org).value, ws.cell(r, c_tab).value
            if not o or not t:
                continue
            if c_m and not is_formula(ws.cell(r, c_m)):
                old_month = old_month or ws.cell(r, c_m).value
                ws.cell(r, c_m).value = MONTHS[month - 1]
            k = (norm(o), tab_key(t))
            if norm(o) == norm(MUNIS) and is_formula(ws.cell(r, c_act)):
                rep['skipped_formula'] += 1
                continue
            setv(ws.cell(r, c_act), idx.get(k, 0.0))
        if old_month and old_month != MONTHS[month - 1]:
            rep['warnings'].append(f'«{ws.title}»: المقدر والمستهدف ومحصل السنة الماضية ما زالت لشهر {old_month}؛ تحتاج ملف المستهدفات (المرحلة الثانية).')
        rep['sheets'].append(ws.title)

    # 3) القطاع
    ws = sheet(S.get('sector', 'القطاع'))
    if ws is not None:
        h = header_map(ws)
        c_org, c_sec, c_tab = h.get(norm('التنظيم الايداري')), h.get(norm('القطاع')), h.get(norm('التبويب'))
        c_b = next((c for k, c in h.items() if 'اجمالي الفواتير' in k), None)
        c_a, c_p, c_t = h.get(norm('فعلي')), h.get(norm('محصل  سنوات سابقة')), h.get(norm('إجمالي المحصل'))
        idx = {(norm(o), norm(s), tab_key(t)): v for (o, s, t), v in agg.sectors.items()}
        for r in range(2, ws.max_row + 1):
            o, s, t = ws.cell(r, c_org).value, ws.cell(r, c_sec).value, ws.cell(r, c_tab).value
            if not o or not s or not t:
                continue
            v = idx.get((norm(o), norm(s), tab_key(t)))
            if v is None:
                rep['unmatched'].append(f'{ws.title}: {o} / {s} / {t}')
                continue
            for c, val in ((c_b, v['billed']), (c_a, v['actual']), (c_p, v['prev_years']), (c_t, v['actual'] + v['prev_years'])):
                if c:
                    setv(ws.cell(r, c), val)
        rep['sheets'].append(ws.title)

    # 4) المفوتر الشهري: صف الشهر المطلوب فقط، والأشهر السابقة نقارنها بس
    ws = sheet(S.get('series', 'المفوتر'))
    if ws is not None:
        h = header_map(ws)
        c_y, c_m = h.get(norm('السنة')), h.get(norm('الشهر'))
        c_b, c_a, c_p = h.get(norm('مفوتر')), h.get(norm('محصل السنة الحالية')), h.get(norm('محصل السنوات السابقة'))
        target_row = None
        last = 1
        for r in range(2, ws.max_row + 1):
            y, m = ws.cell(r, c_y).value, ws.cell(r, c_m).value
            if y is None and m is None:
                continue
            last = r
            if to_num(y) == year and to_num(m) == month:
                target_row = r
            elif to_num(y) == year and 0 < to_num(m) < month:
                s = agg.series.get(int(to_num(m)))
                old = to_num(ws.cell(r, c_b).value) if c_b else 0
                if s and old and abs(s['billed'] - old) / old > 0.005:
                    rep['prev_month_diffs'].append({'month': int(to_num(m)), 'file': round(old, 2), 'now': round(s['billed'], 2)})
        if target_row is None:
            target_row = last + 1
            ws.cell(target_row, c_y).value = year
            ws.cell(target_row, c_m).value = month
        s = agg.series.get(month, {'billed': 0.0, 'actual': 0.0, 'prev_years': 0.0})
        for c, val in ((c_b, s['billed']), (c_a, s['actual']), (c_p, s['prev_years'])):
            if c:
                setv(ws.cell(target_row, c), val)
        rep['sheets'].append(ws.title)

    wb.save(dst)
    return rep


# ---------------------------------------------------------------- summary & checks
def read_targets(template_path):
    """يقرأ المقدر والمستهدف ومحصل السنة الماضية من القالب (للمقارنة فقط)."""
    out = {}
    try:
        wb = openpyxl.load_workbook(template_path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if norm(ws.title) != norm('الايراد حتى تاريخه'):
                continue
            rows = ws.iter_rows(values_only=True)
            hdr = [norm(c) for c in next(rows)]
            ix = {k: hdr.index(norm(k)) for k in ('التنظيم الايداري', 'التبويب', 'مقدر', 'مستهدف', 'محصل السنة الماضية') if norm(k) in hdr}
            for r in rows:
                if not r or r[ix['التنظيم الايداري']] is None:
                    continue
                out[(norm(r[ix['التنظيم الايداري']]), norm(r[ix['التبويب']]))] = {
                    'estimate': to_num(r[ix['مقدر']]) if 'مقدر' in ix else None,
                    'target': to_num(r[ix['مستهدف']]) if 'مستهدف' in ix else None,
                    'last_year': to_num(r[ix['محصل السنة الماضية']]) if 'محصل السنة الماضية' in ix else None}
    except Exception:
        pass
    return out


def build_summary(agg, stats, tpl, year, month, unc_total):
    tg = read_targets(tpl.get('_src', '')) if tpl.get('_src') else {}
    rows = []
    for org in (REGION, CITY, MUNIS):
        for tab, v in sorted(agg.ytd.get(org, {}).items()):
            t = tg.get((norm(org), norm(tab)), {})
            row = {'org': org, 'tab': tab, **{k: round(x, 2) for k, x in v.items()}}
            if t.get('target'):
                row['target'] = round(t['target'], 2)
                row['dev_target'] = round(v['actual'] - t['target'], 2)
                row['dev_target_pct'] = round((v['actual'] - t['target']) / t['target'] * 100, 1)
            if t.get('last_year'):
                row['last_year'] = round(t['last_year'], 2)
                row['growth_pct'] = round((v['actual'] - t['last_year']) / t['last_year'] * 100, 1)
            if v['billed']:
                row['collection_pct'] = round(v['actual'] / v['billed'] * 100, 1)
            rows.append(row)
    return {
        'period': {'year': year, 'month': month, 'month_name': MONTHS[month - 1]},
        'sources': stats,
        'ytd': rows,
        'month': [{'org': o, 'tab': t, 'actual': round(v, 2)} for o in (REGION, CITY, MUNIS) for t, v in sorted(agg.month.get(o, {}).items())],
        'series': [{'month': m, **{k: round(x, 2) for k, x in v.items()}} for m, v in sorted(agg.series.items())],
        'unclassified': {'accounts': len(agg.unclassified), 'amount': round(unc_total, 2),
                         'top': [{'account': a, **u} for a, u in sorted(agg.unclassified.items(), key=lambda kv: -kv[1]['amount'])[:10]]},
        'no_sector_amount': {k: round(v, 2) for k, v in agg.no_sector.items()},
        'template': {k: v for k, v in tpl.items() if not k.startswith('_')},
    }


def salah_checks(summary, agg):
    """تدقيق الأرقام آلياً (شغل Salah): كل فحص يطلع ناجح أو يرفع ملاحظة."""
    checks = []

    def add(name, ok, note):
        checks.append({'check': name, 'ok': bool(ok), 'note': note})

    for f in ('billed', 'actual', 'prev_years'):
        reg = sum(v[f] for v in agg.ytd[REGION].values())
        parts = sum(v[f] for o in (CITY, MUNIS) for v in agg.ytd.get(o, {}).values())
        nulls = sum(v[f] for v in agg.ytd.get(None, {}).values())
        add(f'المنطقة = المدينة + البلديات ({f})', abs(reg - parts - nulls) < 1, f'المنطقة {fmt_m(reg)}، المدينة والبلديات {fmt_m(parts)}، بدون جهة {fmt_m(nulls)} مليون')
    series_b = sum(v['billed'] for v in agg.series.values())
    reg_b = sum(v['billed'] for v in agg.ytd[REGION].values())
    add('مجموع الأشهر = المفوتر حتى تاريخه', abs(series_b - reg_b) < 1, f'{fmt_m(series_b)} مقابل {fmt_m(reg_b)} مليون')
    total = reg_b or 1
    share = summary['unclassified']['amount'] / total * 100
    add('غير المصنف أقل من 0.5٪', share < 0.5, f'{share:.2f}٪ من المفوتر')
    neg = [f"{r['org']}/{r['tab']}" for r in summary['ytd'] if r['billed'] < 0 or r['actual'] < 0]
    add('ما فيه مجاميع سالبة', not neg, '، '.join(neg) or 'سليم')
    add('كل صفوف القالب انربطت', not summary['template']['unmatched'], f"{len(summary['template']['unmatched'])} مفتاح ناقص")
    d = summary['template']['prev_month_diffs']
    add('الأشهر السابقة ما تغيّرت عن الملف السابق', not d, '، '.join(f"شهر {x['month']}: {fmt_m(x['file'])} ← {fmt_m(x['now'])}" for x in d) or 'سليم')
    ok = sum(c['ok'] for c in checks)
    status('salah', 'خلّص تدقيق الأرقام', 'done' if ok == len(checks) else 'alert', detail=f'{ok} من {len(checks)} فحص سليم',
           screen=[('✓ ' if c['ok'] else '✗ ') + c['check'] for c in checks])
    bad = [c for c in checks if not c['ok']]
    if bad:
        chat('salah', 'angel', 'ملاحظات التدقيق: ' + ' | '.join(f"{c['check']}: {c['note']}" for c in bad), kind='return')
    else:
        chat('salah', 'neves', f'دققت الأرقام: {len(checks)} فحص كلها سليمة.', kind='approve')
    return checks


# ---------------------------------------------------------------- numcheck (قفل الأرقام)
_NUM = re.compile(r'([-+−]?\d[\d,]*(?:\.\d+)?)\s*(٪|%|مليار|مليون|ألف|الف)?')


def allowed_values(obj, acc):
    if isinstance(obj, bool):
        return
    if isinstance(obj, (int, float)):
        acc.append(float(obj))
    elif isinstance(obj, dict):
        for v in obj.values():
            allowed_values(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            allowed_values(v, acc)


def numcheck(summary, text, extra=None):
    vals = []
    allowed_values(summary, vals)
    allowed_values(extra or [], vals)
    # الفروق والنسب اللي يحسبها الكاتب من الأرقام نفسها مسموحة إذا كانت مشتقة مباشرة
    bad = []
    for m in _NUM.finditer(text.translate(_DIGITS)):
        raw, unit = m.group(1).replace('−', '-'), m.group(2)
        try:
            n = float(raw.replace(',', ''))
        except ValueError:
            continue
        if unit in ('٪', '%'):
            ok = any(abs(abs(n) - abs(v)) <= 0.051 for v in vals)
        elif unit in ('مليون', 'مليار', 'ألف', 'الف'):
            mul = {'مليون': 1e6, 'مليار': 1e9, 'ألف': 1e3, 'الف': 1e3}[unit]
            dec = len(raw.split('.')[1]) if '.' in raw else 0
            tol = 0.5 * 10 ** -dec * mul + 1
            ok = any(abs(abs(n * mul) - abs(v)) <= tol for v in vals)
        else:
            ok = (n.is_integer() and (0 <= n <= 31 or 1990 <= n <= 2100)) or any(abs(abs(n) - abs(v)) <= 0.51 for v in vals)
        if not ok:
            bad.append(m.group(0).strip())
    return bad


# ---------------------------------------------------------------- inspect
def run_inspect(cfg):
    out = []
    for key, path in cfg['_paths'].items():
        item = {'source': key, 'path': path, 'exists': bool(path and os.path.exists(path))}
        if item['exists']:
            item['size_mb'] = round(os.path.getsize(path) / 1e6, 1)
            try:
                if path.lower().endswith(('.csv', '.txt')):
                    _, rows = open_rows(path)
                    item['sheets'] = [{'name': os.path.basename(path), 'header_candidates': [list(next(rows))]}]
                else:
                    wb = openpyxl.load_workbook(path, read_only=True)
                    item['sheets'] = []
                    for ws in wb.worksheets:
                        heads = []
                        for i, r in enumerate(ws.iter_rows(max_row=8, values_only=True)):
                            texts = [str(c) for c in r if isinstance(c, str)]
                            if len(texts) >= 2:
                                heads.append([c if isinstance(c, str) else None for c in r])
                        item['sheets'].append({'name': ws.title, 'max_row': ws.max_row, 'max_col': ws.max_column,
                                               'header_candidates': [[h for h in hs if h] for hs in heads[:2]]})
                for s in item['sheets']:
                    s['pii_columns'] = sorted({h for hs in s['header_candidates'] for h in hs if h and is_pii(h, cfg)})
            except Exception as e:
                item['error'] = str(e)
        out.append(item)
        emit('inspect', item=item)
    return out


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['inspect', 'monthly', 'numcheck'])
    ap.add_argument('--config')
    ap.add_argument('--year', type=int)
    ap.add_argument('--month', type=int)
    ap.add_argument('--out')
    ap.add_argument('--summary')
    ap.add_argument('--text')
    a = ap.parse_args()
    try:
        if a.cmd == 'inspect':
            run_inspect(load_config(a.config))
        elif a.cmd == 'monthly':
            cfg = load_config(a.config)
            run_monthly(cfg, a.year, a.month, a.out)
        elif a.cmd == 'numcheck':
            s = json.load(open(a.summary, encoding='utf-8'))
            bad = numcheck(s, open(a.text, encoding='utf-8').read())
            emit('numcheck', ok=not bad, unknown=bad)
    except Exception as e:
        emit('error', message=str(e), type=type(e).__name__)
        sys.exit(1)


if __name__ == '__main__':
    main()
